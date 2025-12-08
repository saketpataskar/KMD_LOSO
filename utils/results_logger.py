import pandas as pd
from openpyxl import Workbook, load_workbook
import os

# -------------------------------------------
# Always save Excel logs in project root
# -------------------------------------------
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_FILE = os.path.join(PROJECT_ROOT, "HARTH_Results.xlsx")


def init_excel():
    """Create Excel file with default sheets if it does not exist."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        wb.remove(wb.active)

        # Create main sheets
        for sheet in ["FixedSplit", "KFold", "GroupKFold", "LOSO"]:
            ws = wb.create_sheet(sheet)
            # leave sheet empty — headers added dynamically

        wb.save(EXCEL_FILE)


def append_row(sheet_name, row_dict):
    """Append a row (dictionary) to the specified sheet."""
    init_excel()

    wb = load_workbook(EXCEL_FILE)

    if sheet_name not in wb.sheetnames:
        wb.create_sheet(sheet_name)

    ws = wb[sheet_name]

    # If the sheet is empty → write header first
    if ws.max_row == 1:
        ws.append(list(row_dict.keys()))

    # Write row values
    ws.append([row_dict[k] for k in row_dict])

    wb.save(EXCEL_FILE)
    print(f"[LOGGED] {sheet_name}: {row_dict}")
